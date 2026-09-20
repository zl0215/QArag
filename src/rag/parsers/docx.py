"""DOCX 解析（python-docx，MIT）。

★ 三个必须处理的点：
1. **按 body XML 顺序遍历**，不能用 `doc.paragraphs` —— 后者不含表格，
   且把表格和正文的相对次序丢掉了。
2. **中文 Word 模板的样式名是"标题 1"而不是"Heading 1"** —— 必须做样式名映射，
   否则标题全部退化为普通段落，heading 面包屑直接失效。
3. `.doc` 遗留格式本库不支持（在 base.sniff_mime 已被拦截）。

局限：`doc.paragraphs` 不含页眉页脚/文本框/脚注；`p.text` 会丢掉局部格式（加粗等）。
页眉页脚属于 furniture，本来就要丢弃；脚注如需保留可另接 docx2python。
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from rag.core.logging import get_logger
from rag.parsers.base import finalize
from rag.parsers.tables import table_to_markdown
from rag.schemas.document import DocMeta, DocNode, NodeType, ParsedDocument, make_node_id

logger = get_logger(__name__)

PARSER_NAME = "python-docx"

# 同时覆盖 "Heading 1" / "标题 1" / "h1" / "Heading1"
_HEADING_RE = re.compile(r"^\s*(?:heading|标题|h)\s*([1-9])\s*$", re.IGNORECASE)
# 无编号列表：Word 的 List Paragraph / 列表段落
_LIST_STYLE_RE = re.compile(r"(list\s*paragraph|列表段落|列表)", re.IGNORECASE)


def _heading_level(paragraph: Paragraph) -> int | None:
    """从样式名或 w:outlineLvl 推断标题层级。"""
    style = paragraph.style
    if style is not None:
        for candidate in (style.name, getattr(style, "style_id", None)):
            if not candidate:
                continue
            match = _HEADING_RE.match(str(candidate))
            if match:
                return int(match.group(1))

    # 兜底：直接读大纲级别（某些模板不设样式名，只设 outlineLvl）
    p_pr = paragraph._p.pPr
    if p_pr is not None:
        outline = p_pr.find(qn("w:outlineLvl"))
        if outline is not None:
            value = outline.get(qn("w:val"))
            if value is not None and value.isdigit():
                return min(int(value) + 1, 9)
    return None


def _is_list_item(paragraph: Paragraph) -> bool:
    p_pr = paragraph._p.pPr
    if p_pr is not None and p_pr.find(qn("w:numPr")) is not None:
        return True
    style = paragraph.style
    return bool(style is not None and style.name and _LIST_STYLE_RE.search(style.name))


def _table_rows(table: Table) -> list[list[str]]:
    """展开合并单元格：python-docx 对横向合并的格子会重复返回同一 cell 对象。"""
    rows: list[list[str]] = []
    for row in table.rows:
        cells: list[str] = []
        seen: set[int] = set()
        for cell in row.cells:
            key = id(cell._tc)
            if key in seen:
                cells.append("")          # 横向合并的重复格留空，保持列数对齐
            else:
                seen.add(key)
                cells.append(cell.text.strip())
        rows.append(cells)
    return rows


def parse_docx(path: Path, meta: DocMeta) -> ParsedDocument:
    document = Document(str(path))
    nodes: list[DocNode] = []
    reading_order = 0
    cursor = 0
    heading_stack: list[tuple[int, str]] = []

    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue

            level = _heading_level(paragraph)
            if level is not None:
                while heading_stack and heading_stack[-1][0] >= level:
                    heading_stack.pop()
                heading_stack.append((level, text))
                node_type = NodeType.TITLE
            elif _is_list_item(paragraph):
                node_type = NodeType.LIST
            else:
                node_type = NodeType.PARAGRAPH

            node = DocNode(
                node_id=make_node_id(meta.doc_id, reading_order),
                doc_id=meta.doc_id,
                type=node_type,
                level=level,
                reading_order=reading_order,
                text=text,
                heading_path=" > ".join(t for _, t in heading_stack),
                char_start=cursor,
                char_end=cursor + len(text),
            )
            cursor += len(text) + 2
            nodes.append(node)
            reading_order += 1

        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            rows = _table_rows(table)
            markdown = table_to_markdown(rows)
            if not markdown:
                continue
            node = DocNode(
                node_id=make_node_id(meta.doc_id, reading_order),
                doc_id=meta.doc_id,
                type=NodeType.TABLE,
                reading_order=reading_order,
                text=markdown,
                table_markdown=markdown,
                table_cells=rows,
                heading_path=" > ".join(t for _, t in heading_stack),
                char_start=cursor,
                char_end=cursor + len(markdown),
            )
            cursor += len(markdown) + 2
            nodes.append(node)
            reading_order += 1

    # 页码信息 DOCX 里拿不到（分页由渲染器决定），统一记 0
    meta.parse_warnings.append("DOCX 无固定分页，页码信息不可用")
    return finalize(meta, nodes, parser_name=f"{PARSER_NAME}:docx")
